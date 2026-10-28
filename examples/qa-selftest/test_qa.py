"""Behavioral regression checks for density, native exhibits and PowerPoint XML."""
from __future__ import annotations

import importlib.util
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / 'skills' / 'lyceum' / 'scripts'


def load(name):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / f'{name}.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


qa = load('qa_check')
normalizer = load('normalize_pptx')
NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}


def slide():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    return deck, deck.slides.add_slide(deck.slide_layouts[6])


def text(target, value, x=1, y=2, w=10, h=1):
    shape = target.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = value
    run.font.name, run.font.size = 'Segoe UI', Pt(18)
    return shape


def metrics(target):
    issues = []
    result = qa.academic_metrics(target, 1, issues)
    return result, {issue.kind for issue in issues}


class DensityTests(unittest.TestCase):
    def test_sparse_page_is_flagged(self):
        _, target = slide()
        text(target, 'Only one sentence.')
        result, kinds = metrics(target)
        self.assertEqual(result['charts'], 0)
        self.assertIn('density', kinds)
        self.assertIn('exhibit', kinds)

    def test_title_and_footer_do_not_fill_body(self):
        _, target = slide()
        text(target, 'A long title ' * 12, y=0.4)
        text(target, 'Footer ' * 30, y=7.1)
        result, kinds = metrics(target)
        self.assertEqual(result['body_chars'], 0)
        self.assertIn('density', kinds)

    def test_content_rich_text_still_needs_figure_review(self):
        _, target = slide()
        text(target, 'Evidence and experimental settings. ' * 10, h=3)
        _, kinds = metrics(target)
        self.assertNotIn('density', kinds)
        self.assertIn('exhibit', kinds)

    def test_explicit_discussion_role_is_exempt(self):
        _, target = slide()
        text(target, 'Which experiment should we prioritize?')
        target.notes_slide.notes_text_frame.text = 'Lyceum role: discussion'
        result, kinds = metrics(target)
        self.assertEqual(result['role'], 'discussion')
        self.assertFalse(kinds)

    def test_real_native_chart_is_an_exhibit(self):
        _, target = slide()
        data = CategoryChartData()
        data.categories = ['Control', 'Method']
        data.add_series('Demo', [4, 7])
        target.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2),
                                Inches(8), Inches(4), data)
        result, kinds = metrics(target)
        self.assertEqual(result['charts'], 1)
        self.assertFalse(kinds)

    def test_native_diagram_lines_are_recognized(self):
        _, target = slide()
        for x in (1, 4, 7):
            target.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(3), Inches(2), Inches(1))
        # PptxGenJS uses p:sp with line geometry, rather than p:cxnSp.
        edge = target.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3), Inches(3.5), Inches(1), 0)
        edge._element.find('.//a:prstGeom', NS).set('prst', 'line')
        props = edge._element.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}spPr')
        line = etree.SubElement(props, f"{{{NS['a']}}}ln")
        etree.SubElement(line, f"{{{NS['a']}}}tailEnd", type='triangle')
        result, kinds = metrics(target)
        self.assertTrue(result['diagram_candidate'])
        self.assertFalse(kinds)

    def test_horizontal_rules_do_not_make_a_diagram(self):
        _, target = slide()
        for y in (2, 3, 4):
            text(target, 'A paragraph of text', y=y)
            edge = target.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(y + 0.5), Inches(10), 0)
            edge._element.find('.//a:prstGeom', NS).set('prst', 'line')
        result, kinds = metrics(target)
        self.assertFalse(result['diagram_candidate'])
        self.assertIn('exhibit', kinds)

    def test_empty_boxes_are_not_a_diagram(self):
        _, target = slide()
        for x in (1, 4, 7):
            target.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(3), Inches(2), Inches(1))
        result, kinds = metrics(target)
        self.assertFalse(result['diagram_candidate'])
        self.assertIn('exhibit', kinds)


class CompatibilityTests(unittest.TestCase):
    def test_negative_extents_are_errors(self):
        _, target = slide()
        shape = target.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(2), Inches(1))
        shape._element.find('.//a:ext', NS).set('cy', '-100')
        issues = []
        qa.check_slide(1, target, issues, (13.333, 7.5))
        self.assertTrue(any(i.kind == 'geometry' and i.level == 'ERROR' for i in issues))

    def test_invalid_table_anchor_is_an_error(self):
        _, target = slide()
        tab = target.shapes.add_table(1, 1, Inches(1), Inches(2), Inches(8), Inches(1))
        tab.table.cell(0, 0)._tc.get_or_add_tcPr().set('anchor', 'mid')
        issues = []
        qa.check_slide(1, target, issues, (13.333, 7.5))
        self.assertTrue(any(i.kind == 'xml' for i in issues))

    def duplicate_package(self, path, conflicting=False):
        deck, target = slide()
        text(target, '中文 ABC 123')
        deck.save(path)
        with ZipFile(path) as package:
            members = {name: package.read(name) for name in package.namelist()}
        root = etree.fromstring(members['ppt/slides/slide1.xml'])
        paragraph = root.find('.//a:p', NS)
        first = etree.Element(f"{{{NS['a']}}}pPr", algn='l')
        duplicate = etree.Element(f"{{{NS['a']}}}pPr", algn='r' if conflicting else 'l')
        paragraph.insert(0, first)
        paragraph.append(duplicate)
        members['ppt/slides/slide1.xml'] = etree.tostring(root)
        with ZipFile(path, 'w') as package:
            for name, data in members.items():
                package.writestr(name, data)
        return members

    def test_normalization_preserves_text_and_other_members(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'generated.pptx'
            original = self.duplicate_package(path)
            self.assertEqual(normalizer.normalize(path), 1)
            self.assertEqual(normalizer.normalize(path), 0)
            with ZipFile(path) as package:
                for name, data in original.items():
                    if name != 'ppt/slides/slide1.xml':
                        self.assertEqual(package.read(name), data)
            target = Presentation(path).slides[0]
            self.assertEqual(target.shapes[0].text_frame.text, '中文 ABC 123')
            issues = []
            qa.check_slide(1, target, issues, (13.333, 7.5))
            self.assertFalse([i for i in issues if i.kind == 'xml'])

    def test_conflicting_properties_leave_original_untouched(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'generated.pptx'
            self.duplicate_package(path, conflicting=True)
            before = path.read_bytes()
            with self.assertRaises(ValueError):
                normalizer.normalize(path)
            self.assertEqual(path.read_bytes(), before)

    def test_original_negative_fixture_is_still_rejected(self):
        result = subprocess.run([sys.executable, '-X', 'utf8', str(SCRIPTS / 'qa_check.py'),
                                 str(ROOT / 'examples/qa-selftest/bad_deck.pptx'), '--json'],
                                capture_output=True)
        import json
        report = json.loads(result.stdout.decode('utf-8'))
        self.assertEqual(result.returncode, 1)
        kinds = {item['kind'] for item in report['errors'] + report['warnings']}
        self.assertTrue({'overflow', 'placeholder', 'font', 'offcanvas', 'fontsize', 'overlap'} <= kinds)


if __name__ == '__main__':
    unittest.main()
